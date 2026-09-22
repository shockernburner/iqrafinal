from pathlib import Path
from copy import copy
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET
import calendar
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.workbook.properties import CalcProperties

OUT=Path("outputs/IQRA_Plan_Nov2026_Mar2027_USD.xlsx")
w=openpyxl.load_workbook("attached_assets/IQRA_Yearly_plan_1790054453538.xlsx")
s=w.active
s.title="Monthly Plan"
original=w.copy_worksheet(s); original.title="Original Template"
cache={}
def formula(ws,cell,text,value):
    ws[cell]=text
    cache[(ws.title,cell)]=value
def sheet(name,rows):
    ws=w.create_sheet(name)
    for row in rows: ws.append(row)
    return ws
a=sheet("Assumptions",[
["IQRA operating budget — editable planning inputs","Value","Basis / scope"],
["Peak / average concurrent ratio",4,"Peak online users / 24-hour average; assumption, not measured"],
["Average session minutes",10,"Active sessions, not background tabs"],
["Sessions per daily active user",2,"Assumption"],
["DAU / MAU ratio",.2,"20% of monthly active users visit on an average day"],
["Monthly active-user churn",.10,"Replacement acquisition included"],
["Paid share of gross MAU acquisition",.70,"Remaining 30% organic / partnerships; not guaranteed"],
["Cost per acquired monthly active user USD",3,"Not cost per install; blended geography-dependent assumption"],
["Answered questions per DAU per day",2,"No subscription cap assumed; no revenue assumed"],
["AI USD per answered question",.04,"Planning allowance using current Sonnet setup; not measured bill"],
["Answer latency seconds",20,"Capacity assumption, not benchmark"],
["Cash contingency",.20,"20% of operating expenses; not profit"],
["USD per SGD",.78,"Budget FX assumption, not spot quote"],
["SG GST allowance",.09,"Conservative allowance on OSOME base quote; confirm inclusive/exclusive"],
["OSOME indicative SGD package",3772,"Published starting price; nominee director 1yr, secretary/address; confirm scope"],
["Starting MAU before November",0,"Conservative baseline; replace with measured October MAU"],
["Infrastructure USD per 1,000 monthly answers",2,"Non-AI compute/database/egress allowance, not Replit quote"],
["Infrastructure base monthly USD",500,"Fixed minimum monitoring/hosting allocation"],
["Support monthly USD per 10,000 MAU",150,"Nontechnical support estimate"],
["Accounting/bank admin monthly USD",400,"Allowance; reduce if covered by confirmed OSOME bundle"],
])
g=sheet("Usage and Acquisition",[
["Month","Peak online target","Avg-month peak","Average concurrent","DAU","MAU end target","Gross new MAU","Paid new MAU","Ad spend USD","Monthly answers","AI USD","Peak answer concurrency"],
])
months=["Nov 2026","Dec 2026","Jan 2027","Feb 2027","Mar 2027"]
targets=[1000,2000,4000,8000,16000]
vals=[]
prev=0
for i,(month,target) in enumerate(zip(months,targets),2):
    year=2026 if i<4 else 2027
    mon=[11,12,1,2,3][i-2]
    days=calendar.monthrange(year,mon)[1]
    avgpeak=(prev+target)/2
    avgcon=avgpeak/4; dau=avgcon*1440/10/2
    mau=target/4*1440/10/2/.2
    prevmau=prev/4*1440/10/2/.2
    new=mau-prevmau*.9; paid=new*.7
    q=dau*2*days; ai=q*.04
    peakq=target/10/60/2*2*20
    g.cell(i,1,month);g.cell(i,2,target)
    formula(g,f"C{i}",f"=(B{i}+{('0' if i==2 else 'B'+str(i-1))})/2",avgpeak)
    formula(g,f"D{i}",f"=C{i}/Assumptions!B2",avgcon)
    formula(g,f"E{i}",f"=D{i}*1440/Assumptions!B3/Assumptions!B4",dau)
    formula(g,f"F{i}",f"=B{i}/Assumptions!B2*1440/Assumptions!B3/Assumptions!B4/Assumptions!B5",mau)
    formula(g,f"G{i}",f"=MAX(0,F{i}-{('Assumptions!B16' if i==2 else 'F'+str(i-1))}*(1-Assumptions!B6))",new)
    formula(g,f"H{i}",f"=G{i}*Assumptions!B7",paid)
    formula(g,f"I{i}",f"=H{i}*Assumptions!B8",paid*3)
    formula(g,f"J{i}",f"=E{i}*Assumptions!B9*{days}",q)
    formula(g,f"K{i}",f"=J{i}*Assumptions!B10",ai)
    formula(g,f"L{i}",f"=B{i}/Assumptions!B3/60/Assumptions!B4*Assumptions!B9*Assumptions!B11",peakq)
    vals.append(dict(mau=mau,ads=paid*3,q=q,ai=ai))
    prev=target

l=sheet("Legal and Setup",[
["Item","Nov 2026 USD","Dec","Jan","Feb","Mar","Basis / inclusion"],
["OSOME SG incorporation + foreign-founder compliance",None,0,0,0,0,"Indicative starting bundle; first-year nominee, secretary and registered address. Confirm ACRA fees, accounting scope and tax with OSOME."],
["Legal documents / privacy / PDPA / religious-content review",2000,500,0,0,0,"Professional-services allowance; not an official fee"],
["Bank opening/KYC assistance",300,0,0,0,0,"Allowance, not bank application fee; approval subject to KYC"],
["Stripe registration fee",0,0,0,0,0,"No standard setup fee; technical integration included in Tech"],
["Standard D-U-N-S registration",0,0,0,0,0,"Free standard route; expedited third-party fees excluded"],
["Google Play developer registration",25,0,0,0,0,"Official one-time USD fee"],
["Apple Developer annual membership",99,0,0,0,0,"Official annual USD fee; enrollment subject to verification"],
["Store listing, privacy forms and submission assistance",0,800,400,0,0,"Service allowance; development/testing in Tech; no invented per-app publishing fee"],
["Accounting/tax filings and bank admin allowance",400,400,400,400,400,"Not a tax provision. Reconcile OSOME accounting inclusion before purchase"],
["Total",None,None,None,None,None,"Paid cash outlay; no deposits counted as expenses"],
])
osome=round(3772*.78*1.09,2)
formula(l,"B2","=ROUND(Assumptions!B15*Assumptions!B13*(1+Assumptions!B14),2)",osome)
legal=[]
for col in range(2,7):
    letter=openpyxl.utils.get_column_letter(col)
    formula(l,f"{letter}10","=Assumptions!B20",400)
    value=sum(cache.get((l.title,f"{letter}{r}"),l.cell(r,col).value or 0) for r in range(2,11))
    formula(l,f"{letter}11",f"=SUM({letter}2:{letter}10)",value);legal.append(value)

# Retain the supplied layout, but clear unbudgeted periods and incorrect totals.
for row in range(4,17):
    for col in range(3,17): s.cell(row,col).value=None
s["B3"]="IQRA — November 2026 to March 2027 budget (USD)"
s["B6"]="Peak concurrent users"
s["B7"]="End-month MAU required"
labels={9:"Tech — software development",10:"Promotion — paid advertising",11:"Isnad board / scholar review",12:"HR — nontechnical operations & support",13:"Partnerships / source rights",14:"Legal, company & app-store setup",15:"Reserve — 20% contingency",16:"AI answer generation",17:"Hosting / database / monitoring",18:"App maintenance / QA / security",19:"Ad creative & campaign management",20:"Tools, insurance & devices",21:"Operating subtotal",22:"Total cash budget",23:"Cumulative cash budget"}
for r,label in labels.items():s.cell(r,2,label)
s["H5"]="Nov–Mar total"
dev=[15000,18000,15000,12000,10000]
scholar=[3000,4000,6000,8000,10000]
partner=[1000,2000,3000,5000,8000]
maint=[2000,3000,4500,6000,8000]
tools=[3000,1000,1200,1500,2000]
totals=[]; rowtot={r:0 for r in labels}
for j,v in enumerate(vals):
    col=j+3;c=openpyxl.utils.get_column_letter(col); gr=j+2;lc=openpyxl.utils.get_column_letter(j+2)
    s.cell(4,col,f"Month {j+1}");s.cell(5,col,months[j]);s.cell(6,col,targets[j])
    formula(s,f"{c}7",f"='Usage and Acquisition'!F{gr}",v["mau"])
    numbers={9:dev[j],10:v["ads"],11:scholar[j],12:2000+v["mau"]/10000*150,13:partner[j],14:legal[j],16:v["ai"],17:500+v["q"]/1000*2,18:maint[j],19:1500+v["ads"]*.08,20:tools[j]}
    for r,n in numbers.items():s.cell(r,col,n)
    for r,f in {10:f"='Usage and Acquisition'!I{gr}",12:f"=2000+{c}7/10000*Assumptions!B19",14:f"='Legal and Setup'!{lc}11",16:f"='Usage and Acquisition'!K{gr}",17:f"=Assumptions!B18+'Usage and Acquisition'!J{gr}/1000*Assumptions!B17",19:f"=1500+{c}10*8%"}.items():
        formula(s,f"{c}{r}",f,numbers[r])
    subtotal=sum(numbers.values());reserve=subtotal*.2;total=subtotal+reserve
    formula(s,f"{c}21",f"=SUM({c}9:{c}14,{c}16:{c}20)",subtotal)
    formula(s,f"{c}15",f"={c}21*Assumptions!B12",reserve)
    formula(s,f"{c}22",f"={c}21+{c}15",total)
    totals.append(total)
    formula(s,f"{c}23",f"=SUM($C22:{c}22)",sum(totals))
    numbers.update({15:reserve,21:subtotal,22:total})
    for r,n in numbers.items():rowtot[r]+=n
for r in labels:
    if r!=23:formula(s,f"H{r}",f"=SUM(C{r}:G{r})",rowtot[r])
s["H6"]=16000;s["H7"]=vals[-1]["mau"]
s["H6"].comment=Comment("Ending peak, not sum of concurrent users.","IQRA")
s["H7"].comment=Comment("Ending monthly active audience; not sum of months.","IQRA")
s["B25"]="April–October intentionally not budgeted. Original Template preserves the uploaded sheet unchanged."
s["B26"]="Estimates, not vendor quotes or a capacity guarantee. No revenue, subscriptions or profit assumed."
s["B27"]="Peak targets are month-end; usage costs use a linear within-month ramp. March full run-rate AI cost is higher."
for r in (25,26,27):
    s.merge_cells(start_row=r,start_column=2,end_row=r,end_column=8)
readme=sheet("Read Me",[
["READ FIRST","Planning interpretation and scope"],
["Period","November 2026–March 2027 inferred from uploaded month headings and current date."],
["Target","16,000 peak concurrent online users by March, NOT 16,000 registered users or simultaneous AI requests."],
["Acquisition model","Peak/average ratio 4, 10-minute sessions, 2 sessions/DAU, DAU/MAU 20%. March requires 1.44m MAU and 288k end-state DAU under these assumptions."],
["Advertising","$3 per acquired monthly active user, 70% paid share, 10% monthly MAU churn. CAC and organic share are unvalidated; ads cannot guarantee concurrent attendance."],
["Usage","Two answers/DAU/day. Monthly cost uses average of previous/current peak target, including ramp from zero in November. March average DAU 216k; full run-rate DAU 288k."],
["Capacity","16k peak sessions imply about 533 simultaneously generating answers at 20-second latency under the assumed question rate. Current app capacity is NOT verified; quotas, distributed admission control and load testing required."],
["AI cost","$0.04/answer budget, not measured unit cost. No assumption that all knowledge documents or training rows are sent per request. March full-target run-rate: 17.856m answers / 31 days, $714,240 AI."],
["Tech vs maintenance","Tech: engineering new features, mobile apps, capacity controls and payment integration. Maintenance: separate ongoing bugs, QA, security and on-call. HR excludes these technical roles."],
["Legal scope","First-year incorporation/compliance cash outlay in November; no duplicate nominee/secretary/address charge. Accounting allowance may overlap package: confirm quote and adjust."],
["Bank / capital","Bank approval and Stripe activation are not guaranteed. Restricted minimum balances, paid-up capital, and refundable deposits are assets, not expenses. Hold a separate provisional $2,000 liquidity buffer; bank-dependent and excluded from expense total."],
["Payment fees","Base model assumes no receipts, hence $0 processing fees. If charging users: add Stripe/FX/refund/chargeback fees on web sales OR applicable store commission on store-billed sales; do not double count both on the same sale."],
["Store commissions","Separate from registration. Model 15%–30% of eligible app-store receipts only after confirming program/transaction eligibility. Not assumed included in zero-revenue budget."],
["Taxes","Company profit tax and sales/GST/VAT collection depend on revenue/residency. No profit-tax estimate without revenue. OSOME GST provision is conservative; confirm actual invoice and other vendor taxes."],
["Renewals","Apple annual and SG compliance renewals occur beyond this five-month horizon. Company status and accounting fees can scale with transaction volume."],
["Risk / budget range","20% contingency does not cover arbitrary demand. Adjust CAC, token costs and session behavior before funding. Estimates are geographically blended; high-income-market CAC may be much higher."],
["Scope exclusions","No founder profit/salary unless included in HR; no commercial source-rights quote, physical office, relocation/visa, or unusual regulatory license assumed. Partnership line includes provisional source licensing."],
["Immediate next gate","Validate 30-day cost per answer, retention, paid-acquisition CAC and load-test performance before scaling spend. Nothing in this workbook changes the app or purchases services."],
])
sens=sheet("Sensitivity",[
["Five-month sensitivity","AI cost/answer USD","Paid MAU CAC USD","Cash incl. 20% reserve USD","Interpretation"],
["Lower-cost case",.02,1.5,None,"Not guaranteed; all traffic and fixed assumptions unchanged"],
["Base",.04,3,None,"Working budget"],
["Higher-cost case",.06,6,None,"Not a worst-case ceiling"],
])
base=sum(totals);qsum=sum(v["q"] for v in vals);paidsum=sum(v["ads"]/3 for v in vals)
for r,ai,cac in [(2,.02,1.5),(3,.04,3),(4,.06,6)]:
    value=base+(qsum*(ai-.04)+paidsum*(cac-3)*1.08)*1.2
    formula(sens,f"D{r}",f"='Monthly Plan'!H22+(SUM('Usage and Acquisition'!J2:J6)*(B{r}-Assumptions!B10)+SUM('Usage and Acquisition'!H2:H6)*(C{r}-Assumptions!B8)*1.08)*(1+Assumptions!B12)",value)
src=sheet("Sources",[
["Source (checked 22 Sep 2026)","URL","Use / limitation"],
["OSOME nominee director packages","https://osome.com/sg/nominee-director/","Starting SGD 3,772 bundle; binding quote and tax inclusions required"],
["Apple membership","https://developer.apple.com/programs/whats-included/","USD 99 annual; commissions depend on eligibility"],
["Apple D-U-N-S","https://developer.apple.com/support/D-U-N-S/","Standard organization number free"],
["Google Play registration","https://support.google.com/googleplay/android-developer/answer/6112435","USD 25 one-time"],
["Stripe Singapore","https://stripe.com/sg/pricing","No setup/monthly fee; transaction charges depend on method/currency"],
["Replit AI integration","https://docs.replit.com/features/integrations/replit-ai-integrations","Provider public API prices; actual consumption billed"],
["Replit usage billing","https://docs.replit.com/billing/about-usage-based-billing","Compute/database/transfer depend on usage; workbook uses allowance, not quote"],
["Claude Sonnet 4.6","https://www.anthropic.com/news/claude-sonnet-4-6","Prior pricing basis $3 input/$15 output per million; cost per answer is assumption"],
["All other costs","Management planning estimates","Not supplier quotes: geography, hiring model and commercial rights change costs"],
])
for ws in w:
    if ws==original:continue
    ws.freeze_panes="C6" if ws==s else "B2"
    ws.sheet_view.showGridLines=False
    ws.column_dimensions["A"].width=3 if ws==s else 46
    for col in range(2,ws.max_column+1):ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width=22
    if ws==s:ws.column_dimensions["B"].width=43
    elif ws in (readme,src):ws.column_dimensions["B"].width=100
    if ws==a:ws.column_dimensions["C"].width=95
    if ws==l:ws.column_dimensions["G"].width=90
    for row in ws:
        for cell in row:
            cell.alignment=Alignment(vertical="top",wrap_text=True)
            cell.font=Font(name="Calibri",size=11,color="163B30")
            if isinstance(cell.value,(int,float)) or (isinstance(cell.value,str) and cell.value.startswith("=")):
                cell.number_format='#,##0.00;[Red](#,##0.00);–'
                cell.font=Font(name="Calibri",size=11,color="1663A6" if cell.data_type!="f" else "163B30")
        ws.row_dimensions[row[0].row].height=45 if ws in (readme,src,a,l) else 32
    header=5 if ws==s else 1
    for cell in ws[header]:
        cell.fill=PatternFill("solid",fgColor="123D30");cell.font=Font(name="Calibri",bold=True,color="FFFFFF")
    ws.print_options.horizontalCentered=True
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
for row in (15,21,22,23):
    for cell in s[row][1:8]:
        cell.fill=PatternFill("solid",fgColor="E2EEDC");cell.font=Font(bold=True,color="123D30")
for r in (25,26,27):s.row_dimensions[r].height=45
for r in range(2,readme.max_row+1):readme.row_dimensions[r].height=66
for r in range(2,l.max_row+1):l.row_dimensions[r].height=60
for r in (5,6,7,12,14):
    a.cell(r,2).number_format="0%"
for r in range(9,24):
    for c in range(3,9):s.cell(r,c).number_format='"$"#,##0.00'
s.print_area="B3:H27"
w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
w.save(OUT)
# Populate formula caches so previews show figures without requiring Excel recalculation.
ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
with ZipFile(OUT) as z:entries={n:z.read(n) for n in z.namelist()}
for idx,ws in enumerate(w,1):
    path=f"xl/worksheets/sheet{idx}.xml";root=ET.fromstring(entries[path])
    for cell in root.findall(".//m:c",ns):
        key=(ws.title,cell.attrib["r"])
        if key in cache:
            v=cell.find("m:v",ns)
            if v is None:v=ET.SubElement(cell,"{"+ns["m"]+"}v")
            v.text=str(cache[key])
    entries[path]=ET.tostring(root,encoding="utf-8")
with ZipFile(OUT,"w",ZIP_DEFLATED) as z:
    for name,data in entries.items():z.writestr(name,data)
check=openpyxl.load_workbook(OUT,data_only=True)
assert abs(check["Monthly Plan"]["H22"].value-sum(totals))<.01
assert check["Usage and Acquisition"]["F6"].value==1440000
assert len(cache)==sum(1 for ws in w if ws!=original for row in ws for cell in row if cell.data_type=="f")
print("Monthly cash:",[round(t,2) for t in totals])
print("Total cash:",round(sum(totals),2))
print("Ads:",sum(v["ads"] for v in vals),"AI:",sum(v["ai"] for v in vals))
print("Sensitivity:",[check["Sensitivity"].cell(r,4).value for r in range(2,5)])
print("Validated:",OUT)